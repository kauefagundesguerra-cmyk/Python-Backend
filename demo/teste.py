import pandas as pd 
from openpyxl import load_workbook
from copy import copy
from pathlib import Path
def save_format(origin_file,edited_file,name_saved,filetype=""):
  match filetype:
    case "excel":
      if ".xlsx" not in origin_file and name_saved:
        name_saved = name_saved + ".xlsx"
        origin_name = origin_file + ".xlsx"
      wb_origin = load_workbook(origin_name)
      ws_origin = wb_origin.active
      for line in ws_origin:
        for cel in line:
          font = copy(cel.font)
          fill = copy(cel.fill)
          border = copy(cel.border)
          alignment = copy(cel.alignment)
          number_format = cel.number_format
      edited = pd.read_excel(edited_file)
      edited.to_excel("Não Interaja / No Touch.xlsx")
      wb_none = load_workbook("Não Interaja / No Touch.xlsx")
      ws_none= wb_none.active
      for line in ws_none:
        for cel in line:
          cel.font = font
          cel.fill = fill
          cel.border = border
          cel.alignment = alignment
          cel.number_format = number_format
      origin_file = Path(origin_name)
      if origin_file.exists():
        origin_file.unlink()
      return wb_none.save(name_saved)
save_format("/home/kaue1992/Área de trabalho/openeasy/demo/tabela_teste_formatada.xlsx","/home/kaue1992/Área de trabalho/openeasy/demo/tabela_teste_analise_teste.xlsx","analise_teste",filetype="excel")

    